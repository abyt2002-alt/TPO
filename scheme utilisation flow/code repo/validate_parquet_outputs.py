from __future__ import annotations

import json
import re
from pathlib import Path

import pandas as pd
import pyarrow.parquet as pq
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[2]
RAW_DIR = PROJECT_ROOT / "final_presentation" / "Raw files"
OUTPUT_DIR = Path(__file__).resolve().parents[1] / "output"


def safe_stem(name: str) -> str:
    text = Path(name).stem.strip().lower()
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return text.strip("_") or "scheme_utilisation"


def main() -> int:
    raw_files = sorted(p for p in RAW_DIR.glob("*.xlsx") if not p.name.startswith("~$"))
    parquet_files = sorted(
        path for path in OUTPUT_DIR.glob("*.parquet")
        if not path.name.startswith("step2_")
    )

    checks = []
    errors = []
    raw_total_rows = 0
    parquet_total_rows = 0

    for raw_path in raw_files:
        parquet_path = OUTPUT_DIR / f"{safe_stem(raw_path.name)}.parquet"
        if not parquet_path.exists():
            errors.append({"file": raw_path.name, "error": "missing parquet"})
            continue

        wb = load_workbook(raw_path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        raw_rows = max(ws.max_row - 1, 0)
        raw_cols = ws.max_column
        wb.close()

        parquet_file = pq.ParquetFile(parquet_path)
        parquet_rows = int(parquet_file.metadata.num_rows)
        parquet_cols = int(len(parquet_file.schema_arrow.names))
        parquet_total_rows += parquet_rows
        raw_total_rows += raw_rows

        edge_cols = [
            "source_file",
            "source_row_number",
            "Month",
            "Bill No",
            "Outlet_Id",
            "Invoice qty. pieces",
            "Scheme_discount",
            "%_Scheme",
            "Sale_Value",
        ]
        edge_cols = [c for c in edge_cols if c in parquet_file.schema_arrow.names]
        edge_df = pq.read_table(parquet_path, columns=edge_cols).to_pandas()
        first_row = edge_df.iloc[0].to_dict() if len(edge_df) else {}
        last_row = edge_df.iloc[-1].to_dict() if len(edge_df) else {}

        numeric_cols = [
            "Invoice qty. pieces",
            "Scheme_discount",
            "Sale_Value",
            "Distributor_Sale_Value",
            "Sale_Return_Qty",
            "Sale_Return_Amount",
            "Sale_Return_Discount",
            "Free_quantity",
        ]
        numeric_cols = [c for c in numeric_cols if c in parquet_file.schema_arrow.names]
        numeric_df = pq.read_table(parquet_path, columns=numeric_cols).to_pandas()
        sums = {
            c: round(float(pd.to_numeric(numeric_df[c], errors="coerce").fillna(0).sum()), 4)
            for c in numeric_cols
        }

        checks.append({
            "raw_file": raw_path.name,
            "parquet_file": parquet_path.name,
            "raw_rows": raw_rows,
            "parquet_rows": parquet_rows,
            "row_match": raw_rows == parquet_rows,
            "raw_cols": raw_cols,
            "parquet_cols": parquet_cols,
            "expected_parquet_cols": raw_cols + 3,
            "col_match": parquet_cols == raw_cols + 3,
            "source_file_match": str(first_row.get("source_file")) == raw_path.name,
            "first_source_row_number": int(first_row.get("source_row_number", 0) or 0),
            "last_source_row_number": int(last_row.get("source_row_number", 0) or 0),
            "source_row_range_match": (
                int(first_row.get("source_row_number", 0) or 0) == 2
                and int(last_row.get("source_row_number", 0) or 0) == raw_rows + 1
            ),
            **sums,
        })

    summary = {
        "raw_file_count": len(raw_files),
        "parquet_file_count": len(parquet_files),
        "raw_total_rows": raw_total_rows,
        "parquet_total_rows": parquet_total_rows,
        "all_row_counts_match": all(row["row_match"] for row in checks) and not errors,
        "all_column_counts_match": all(row["col_match"] for row in checks) and not errors,
        "all_source_file_names_match": all(row["source_file_match"] for row in checks) and not errors,
        "all_source_row_ranges_match": all(row["source_row_range_match"] for row in checks) and not errors,
        "errors": errors,
        "checks": checks,
    }

    report_path = OUTPUT_DIR / "validation_report.json"
    report_path.write_text(json.dumps(summary, indent=2), encoding="utf-8")

    print(json.dumps({
        k: v
        for k, v in summary.items()
        if k not in {"checks"}
    }, indent=2))
    print(f"validation_report={report_path}")
    return 0 if not errors and summary["all_row_counts_match"] and summary["all_column_counts_match"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
